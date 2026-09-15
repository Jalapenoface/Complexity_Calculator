from __future__ import annotations

import io
import json
import re
from html import unescape
from datetime import datetime
from pathlib import Path
from collections import Counter

from flask import Flask, abort, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


BASE_DIR = Path(__file__).resolve().parent
app = Flask(__name__)

CATEGORY_LABELS = {
    "assessments": "Assessments", "forums": "Forums", "quizzes": "Quizzes",
    "grading-tools": "Grading Tools", "gradebook": "Gradebook",
    "manual-intervention": "Manual Intervention Components", "interactivity": "Interactivity & Peer Work",
    "restrictions": "Restrictions", "groups": "Groups Complexity",
    "3rd-party": "3rd Party Tools or Links", "plugins": "Special Plugins", "instructor": "Instructional Team Factors",
    "instructor-access": "Instructional Team Platform Access", "other-factors": "Other Factors",
}


def load_allowed_components():
    source = (BASE_DIR / "templates" / "index.html").read_text(encoding="utf-8")
    pattern = re.compile(
        r'<input[^>]+data-points="([^"]+)"[^>]+data-category="([^"]+)"[^>]*>'
        r'.*?<span class="flex-1">(.*?)</span>', re.DOTALL
    )
    return Counter((category, re.sub(r"<[^>]+>", "", unescape(label)).strip(), float(points))
                   for points, category, label in pattern.findall(source))


ALLOWED_COMPONENTS = load_allowed_components()


def safe_number(value, default=0.0):
    try:
        number = float(value)
        return number if -1000 <= number <= 1000 else default
    except (TypeError, ValueError):
        return default


def clean_text(value, limit=2000):
    return str(value or "").strip()[:limit]


def safe_quantity(value):
    try:
        quantity = int(value)
        return quantity if 0 <= quantity <= 10000 else 0
    except (TypeError, ValueError):
        return 0


def calculate(payload):
    meta = payload.get("meta") or {}
    class_cap = max(0, int(safe_number(meta.get("class_cap"))))
    cap_points = 2 if class_cap >= 1000 else 1 if class_cap >= 600 else 0
    rows, breakdown = [], {key: 0.0 for key in CATEGORY_LABELS}

    seen = Counter()
    for item in payload.get("selected") or []:
        category = clean_text(item.get("category"), 50)
        if category not in breakdown:
            continue
        points = safe_number(item.get("points"))
        label = clean_text(item.get("label"), 200) or "Selected component"
        component = (category, label, points)
        if seen[component] >= ALLOWED_COMPONENTS[component]:
            continue
        seen[component] += 1
        quantity = safe_quantity(item.get("quantity"))
        calculated_points = points * quantity
        breakdown[category] += calculated_points
        rows.append({"category": CATEGORY_LABELS[category], "item": label, "quantity": quantity,
                     "unit_points": points, "points": calculated_points})

    custom = payload.get("custom") or {}
    for category, entry in custom.items():
        if category not in breakdown or not isinstance(entry, dict):
            continue
        description = clean_text(entry.get("description"), 300)
        points = safe_number(entry.get("points"))
        quantity = safe_quantity(entry.get("quantity"))
        if description or points:
            calculated_points = points * quantity
            breakdown[category] += calculated_points
            rows.append({"category": CATEGORY_LABELS[category], "item": description or "Other complexity",
                         "quantity": quantity, "unit_points": points, "points": calculated_points})

    gamification_points = max(0, safe_number(payload.get("gamification_points")))
    gamification_quantity = safe_quantity(payload.get("gamification_quantity"))
    if gamification_points:
        calculated_points = gamification_points * gamification_quantity
        breakdown["manual-intervention"] += calculated_points
        rows.append({"category": CATEGORY_LABELS["manual-intervention"],
                     "item": "Gamification or interactivity", "quantity": gamification_quantity,
                     "unit_points": gamification_points, "points": calculated_points})

    total = cap_points + sum(breakdown.values())
    return {
        "meta": {
            "class_code": clean_text(meta.get("class_code"), 100), "class_cap": class_cap,
            "filled_by": clean_text(meta.get("filled_by"), 150), "date": clean_text(meta.get("date"), 30),
            "semester": clean_text(meta.get("semester"), 30).title(), "year": clean_text(meta.get("year"), 4),
        },
        "rows": rows, "breakdown": breakdown, "cap_points": cap_points, "total": total,
        "context": clean_text(payload.get("context")), "notes": clean_text(payload.get("notes")),
    }


def display_number(value):
    return str(int(value)) if float(value).is_integer() else f"{value:g}"


def file_stem(report):
    raw = report["meta"]["class_code"] or "course"
    return re.sub(r"[^A-Za-z0-9_-]+", "_", raw).strip("_") or "course"


@app.get("/")
def index():
    return render_template("index.html")


@app.post("/export/<kind>")
def export(kind):
    if kind not in {"pdf", "xlsx"}:
        abort(404)
    try:
        payload = json.loads(request.form.get("payload", "{}"))
    except json.JSONDecodeError:
        abort(400, "Invalid report data")
    report = calculate(payload)
    return export_pdf(report) if kind == "pdf" else export_xlsx(report)


def export_pdf(report):
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter, rightMargin=.55*inch, leftMargin=.55*inch,
                            topMargin=.55*inch, bottomMargin=.55*inch)
    styles = getSampleStyleSheet()
    styles["Title"].textColor = colors.HexColor("#4338CA")
    story = [Paragraph("ICC Course Operational Complexity", styles["Title"]), Spacer(1, 10)]
    m = report["meta"]
    metadata = [
        ["Class Code", m["class_code"] or "—", "Class Cap", str(m["class_cap"])],
        ["Filled by", m["filled_by"] or "—", "Date", m["date"] or "—"],
        ["Semester", m["semester"] or "—", "Year", m["year"] or "—"],
    ]
    meta_table = Table(metadata, colWidths=[.85*inch, 2.15*inch, .75*inch, 2.4*inch])
    meta_table.setStyle(TableStyle([("BACKGROUND", (0,0), (0,-1), colors.HexColor("#E0E7FF")),
                                    ("BACKGROUND", (2,0), (2,-1), colors.HexColor("#E0E7FF")),
                                    ("GRID", (0,0), (-1,-1), .35, colors.HexColor("#CBD5E1")),
                                    ("FONTNAME", (0,0), (-1,-1), "Helvetica"),
                                    ("FONTNAME", (0,0), (0,-1), "Helvetica-Bold"),
                                    ("FONTNAME", (2,0), (2,-1), "Helvetica-Bold"),
                                    ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("PADDING", (0,0), (-1,-1), 6)]))
    story += [meta_table, Spacer(1, 14), Paragraph(f"Total Score: {display_number(report['total'])}", styles["Heading1"]), Spacer(1, 8)]

    detail = [["Category", "Component", "Qty", "Unit pts", "Total pts"],
              ["Enrollment", f"Class cap: {m['class_cap']}", "1", display_number(report["cap_points"]), display_number(report["cap_points"])]]
    detail += [[r["category"], Paragraph(r["item"], styles["BodyText"]), str(r["quantity"]),
                display_number(r["unit_points"]), display_number(r["points"])] for r in report["rows"]]
    if len(detail) == 2 and not report["rows"]:
        detail.append(["—", "No components selected", "—", "0", "0"])
    table = Table(detail, repeatRows=1, colWidths=[1.25*inch, 3.35*inch, .45*inch, .7*inch, .8*inch])
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#4338CA")),
                               ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
                               ("GRID", (0,0), (-1,-1), .35, colors.HexColor("#CBD5E1")),
                               ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
                               ("ALIGN", (2,1), (-1,-1), "RIGHT"), ("VALIGN", (0,0), (-1,-1), "TOP"),
                               ("PADDING", (0,0), (-1,-1), 5)]))
    story += [table, Spacer(1, 14), Paragraph("Score Breakdown", styles["Heading2"])]
    breakdown = [["Enrollment", display_number(report["cap_points"])]] + [
        [CATEGORY_LABELS[k], display_number(v)] for k, v in report["breakdown"].items()
    ] + [["TOTAL", display_number(report["total"])]]
    bt = Table(breakdown, colWidths=[3.2*inch, 1*inch])
    bt.setStyle(TableStyle([("GRID", (0,0), (-1,-1), .35, colors.HexColor("#CBD5E1")),
                            ("BACKGROUND", (0,-1), (-1,-1), colors.HexColor("#E0E7FF")),
                            ("FONTNAME", (0,-1), (-1,-1), "Helvetica-Bold"),
                            ("ALIGN", (1,0), (1,-1), "RIGHT"), ("PADDING", (0,0), (-1,-1), 5)]))
    story.append(bt)
    if report["context"] or report["notes"]:
        story += [Spacer(1, 14), Paragraph("Notes", styles["Heading2"])]
        if report["context"]:
            story += [Paragraph("Changes this semester", styles["Heading3"]), Paragraph(report["context"].replace("\n", "<br/>"), styles["BodyText"])]
        if report["notes"]:
            story += [Paragraph("Final notes", styles["Heading3"]), Paragraph(report["notes"].replace("\n", "<br/>"), styles["BodyText"])]
    doc.build(story)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=f"{file_stem(report)}_complexity_report.pdf", mimetype="application/pdf")


def export_xlsx(report):
    wb = Workbook()
    ws = wb.active
    ws.title = "Complexity Report"
    navy, indigo, pale = "0F172A", "4F46E5", "E0E7FF"
    ws.merge_cells("A1:C1")
    ws["A1"] = "ICC Course Operational Complexity"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=indigo)
    ws["A1"].alignment = Alignment(horizontal="center")
    m = report["meta"]
    info = [("Class Code", m["class_code"]), ("Class Cap", m["class_cap"]), ("Filled by", m["filled_by"]),
            ("Date", m["date"]), ("Semester", m["semester"]), ("Year", m["year"]), ("Total Score", report["total"])]
    row = 3
    for label, value in info:
        ws.cell(row, 1, label).font = Font(bold=True, color=navy)
        ws.cell(row, 2, value)
        row += 1
    row += 1
    for col, value in enumerate(["Category", "Component", "Quantity", "Unit Points", "Calculated Points"], 1):
        c = ws.cell(row, col, value); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=indigo)
    row += 1
    ws.append(["Enrollment", f"Class cap: {m['class_cap']}", 1, report["cap_points"], report["cap_points"]])
    row += 1
    for item in report["rows"]:
        ws.append([item["category"], item["item"], item["quantity"], item["unit_points"], item["points"]]); row += 1
    row += 1
    ws.cell(row, 1, "Score Breakdown").font = Font(size=14, bold=True, color=indigo)
    row += 1
    for label, points in [("Enrollment", report["cap_points"])] + [(CATEGORY_LABELS[k], v) for k, v in report["breakdown"].items()]:
        ws.cell(row, 1, label); ws.cell(row, 5, points); row += 1
    ws.cell(row, 1, "TOTAL").font = Font(bold=True); ws.cell(row, 5, report["total"]).font = Font(bold=True)
    for cell in ws[row]: cell.fill = PatternFill("solid", fgColor=pale)
    row += 2
    ws.cell(row, 1, "Changes This Semester").font = Font(bold=True); ws.cell(row, 2, report["context"]); row += 1
    ws.cell(row, 1, "Final Notes").font = Font(bold=True); ws.cell(row, 2, report["notes"])
    ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 12; ws.column_dimensions["D"].width = 14; ws.column_dimensions["E"].width = 18
    for cells in ws.iter_rows():
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A12"
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, as_attachment=True, download_name=f"{file_stem(report)}_complexity_report.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    app.run(debug=True)
